"""Shared construction and workbook helpers for municipality result tables."""

from __future__ import annotations

from pathlib import Path
from math import isfinite
from typing import Sequence

import geopandas as gpd
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
GRID_SOURCE = ROOT / "data/processed/kumamoto_grid_exposure_estimates_preprocessed.parquet"
POWER_SOURCE = ROOT / "data/processed/kumamoto_emergency_cooling_power_scenarios_preprocessed.parquet"
MORTALITY_SOURCE = ROOT / "data/processed/kumamoto_cooling_loss_mortality_scenario_preprocessed.parquet"
MAP_CRS = "EPSG:6670"

HOUSING_COLUMNS = [
    "Expected Functionally Lost Residences",
    "Expected Functionally Lost Residences Lower Bound",
    "Expected Functionally Lost Residences Upper Bound",
    "Estimated Affected Population Age 65+",
    "Estimated Affected Population Age 65+ Lower Bound",
    "Estimated Affected Population Age 65+ Upper Bound",
]
HEALTH_COLUMNS = [
    "No-Placement Unprotected Older-Person-Days",
    "Expected High-Heat Scenario Days",
    "Incremental Cooling-Loss-Related Excess Deaths",
    "Cooling-Loss Relative 30-Day Mortality Burden Increase %",
    "Cooling-Loss Relative 30-Day Mortality Burden Increase % Lower 95% CI",
    "Cooling-Loss Relative 30-Day Mortality Burden Increase % Upper 95% CI",
]


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def assign_municipalities(
    frame: gpd.GeoDataFrame, municipalities: gpd.GeoDataFrame
) -> gpd.GeoDataFrame:
    projected = frame.to_crs(MAP_CRS).copy()
    points = gpd.GeoDataFrame(
        {"_row_id": np.arange(len(projected))},
        geometry=projected.geometry.representative_point(),
        crs=MAP_CRS,
    )
    joined = gpd.sjoin(
        points,
        municipalities[["Municipality Code", "geometry"]],
        how="left",
        predicate="within",
    ).sort_values("_row_id")
    require(len(joined) == len(projected), "Municipality join changed row count")
    unmatched = joined["Municipality Code"].isna()
    if unmatched.any():
        nearest = gpd.sjoin_nearest(
            joined.loc[unmatched, ["_row_id", "geometry"]],
            municipalities[["Municipality Code", "geometry"]],
            how="left",
            max_distance=100,
            distance_col="Municipality Assignment Distance m",
        )
        nearest = (
            nearest.sort_values(["_row_id", "Municipality Assignment Distance m"])
            .drop_duplicates("_row_id")
            .set_index("_row_id")
        )
        require(nearest["Municipality Code"].notna().all(), "Unassigned grid rows")
        joined = joined.set_index("_row_id")
        joined.loc[nearest.index, "Municipality Code"] = nearest["Municipality Code"]
        joined = joined.reset_index().sort_values("_row_id")
    require(joined["Municipality Code"].notna().all(), "Unassigned grid rows")
    projected["Municipality Code"] = joined["Municipality Code"].astype("string").to_numpy()
    return projected


def build_master_summary() -> pd.DataFrame:
    for source in (GRID_SOURCE, POWER_SOURCE, MORTALITY_SOURCE):
        require(source.is_file(), f"Missing source: {source}")
    mortality = gpd.read_parquet(MORTALITY_SOURCE).to_crs(MAP_CRS)
    mortality["Municipality Code"] = mortality["Municipality Code"].astype("string")
    require(len(mortality) == 45, "Expected 45 municipality rows")
    require(mortality["Municipality Code"].is_unique, "Municipality codes are not unique")
    municipalities = mortality[["Municipality Code", "Municipality", "geometry"]].copy()

    grids = assign_municipalities(gpd.read_parquet(GRID_SOURCE), municipalities)
    require(set(HOUSING_COLUMNS).issubset(grids.columns), "Housing variables are missing")
    housing = grids.groupby("Municipality Code", as_index=False)[HOUSING_COLUMNS].sum()
    for column in HOUSING_COLUMNS:
        require(np.isclose(housing[column].sum(), grids[column].sum()), f"Changed total: {column}")

    power = gpd.read_parquet(POWER_SOURCE)
    require(set(power["Power Demand Scenario"].unique()) == {"Low", "Central", "High"}, "Unexpected power scenarios")
    central = power.loc[power["Power Demand Scenario"].eq("Central")].copy()
    central["Municipality Code"] = "43" + central["Municipality Code"].astype("string").str.zfill(3)
    wards = {"43101", "43102", "43103", "43104", "43105"}
    central.loc[central["Municipality Code"].isin(wards), "Municipality Code"] = "43100"
    require(set(central["Municipality Code"]) == set(municipalities["Municipality Code"]), "Power geography mismatch")
    source_power = [
        "Required Peak Cooling Electric Power kW",
        "Required Daily Cooling Electricity kWh",
    ]
    power_summary = central.groupby("Municipality Code", as_index=False)[source_power].sum()
    power_summary = power_summary.rename(
        columns={
            source_power[0]: "Required Peak Cooling Electric Power kW (Central)",
            source_power[1]: "Required Daily Cooling Electricity kWh (Central)",
        }
    )

    health = pd.DataFrame(mortality.drop(columns="geometry"))[
        ["Municipality Code", "Municipality", *HEALTH_COLUMNS]
    ]
    master = health.merge(housing, on="Municipality Code", validate="one_to_one").merge(
        power_summary, on="Municipality Code", validate="one_to_one"
    )
    require(len(master) == 45 and master.notna().all().all(), "Incomplete master summary")
    require(
        (master["Expected Functionally Lost Residences Lower Bound"] <= master["Expected Functionally Lost Residences"]).all()
        and (master["Expected Functionally Lost Residences"] <= master["Expected Functionally Lost Residences Upper Bound"]).all(),
        "Housing scenario range is invalid",
    )
    require(
        (master["Estimated Affected Population Age 65+ Lower Bound"] <= master["Estimated Affected Population Age 65+"]).all()
        and (master["Estimated Affected Population Age 65+"] <= master["Estimated Affected Population Age 65+ Upper Bound"]).all(),
        "Older-person scenario range is invalid",
    )
    return master


def _visible_number(value: float, decimals: int = 1, thousands: bool = True) -> str:
    """Format a value compactly without hiding a supported nonzero as zero."""
    require(isfinite(float(value)), "Cannot format a non-finite result")
    visible_decimals = decimals
    while value != 0 and round(float(value), visible_decimals) == 0:
        visible_decimals += 1
    grouping = "," if thousands else ""
    return f"{value:{grouping}.{visible_decimals}f}"


def scenario_range(lower: pd.Series, upper: pd.Series) -> pd.Series:
    return lower.map(_visible_number) + "–" + upper.map(_visible_number)


def _nonzero_number_format(value: float, base_format: str) -> str:
    """Increase precision only when the base format would render nonzero as zero."""
    if not isinstance(value, (int, float, np.integer, np.floating)):
        return base_format
    if not isfinite(float(value)) or value == 0 or "." not in base_format:
        return base_format
    decimals = len(base_format.rsplit(".", 1)[1])
    if round(float(value), decimals) != 0:
        return base_format
    while round(float(value), decimals) == 0:
        decimals += 1
    prefix = "#,##0" if "," in base_format else "0"
    return prefix + "." + ("0" * decimals)


def write_result_workbook(
    table: pd.DataFrame,
    output: Path,
    title: str,
    definitions: Sequence[tuple[str, str, str, str]],
    formats: dict[int, str],
    heatmap_columns: Sequence[int],
    footer: str,
) -> None:
    require(table.shape == (45, 6), f"Expected 45 x 6 table, got {table.shape}")
    require(table.notna().all().all(), "Result table contains missing values")
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Municipal Summary"
    sheet.sheet_view.showGridLines = False
    for column_index, header in enumerate(table.columns, start=1):
        sheet.cell(1, column_index, header)
    for row_index, row in enumerate(table.itertuples(index=False, name=None), start=2):
        for column_index, value in enumerate(row, start=1):
            sheet.cell(row_index, column_index, value)

    header_fill = PatternFill("solid", fgColor="17365D")
    thin_gray = Side(style="thin", color="D9E1F2")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 54
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=thin_gray)
            cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "right", vertical="center")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFD")
    for column_index, number_format in formats.items():
        for row_index in range(2, 47):
            cell = sheet.cell(row_index, column_index)
            cell.number_format = _nonzero_number_format(cell.value, number_format)
    for column_index, width in enumerate([24, 23, 24, 24, 25, 25], start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = "A1:F46"
    for column_index in heatmap_columns:
        letter = get_column_letter(column_index)
        sheet.conditional_formatting.add(
            f"{letter}2:{letter}46",
            ColorScaleRule(
                start_type="min", start_color="FFF2CC",
                mid_type="percentile", mid_value=50, mid_color="F4B183",
                end_type="max", end_color="C65911",
            ),
        )
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_title_rows = "1:1"
    sheet.print_area = "A1:F46"
    sheet.oddFooter.center.text = footer
    sheet.oddFooter.right.text = "Page &P of &N"

    definition_sheet = workbook.create_sheet("Definitions")
    definition_sheet.sheet_view.showGridLines = False
    for column_index, header in enumerate(["Column", "Unit", "Definition", "Construction / Limit"], start=1):
        cell = definition_sheet.cell(1, column_index, header)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index, values in enumerate(definitions, start=2):
        for column_index, value in enumerate(values, start=1):
            cell = definition_sheet.cell(row_index, column_index, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        if row_index % 2 == 0:
            for cell in definition_sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor="F3F6FA")
    for column_index, width in enumerate([50, 20, 66, 76], start=1):
        definition_sheet.column_dimensions[get_column_letter(column_index)].width = width
    for row_index in range(2, definition_sheet.max_row + 1):
        definition_sheet.row_dimensions[row_index].height = 44
    definition_sheet.freeze_panes = "A2"
    definition_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    definition_sheet.page_setup.orientation = "landscape"
    definition_sheet.page_setup.fitToWidth = 1
    definition_sheet.page_setup.fitToHeight = 1

    workbook.properties.title = title
    workbook.properties.subject = "2026 Kumamoto earthquake scenario planning"
    workbook.properties.creator = "MiliFrame KE01"
    workbook.save(output)
    validate_workbook(output, list(table.columns))


def validate_workbook(output: Path, headers: list[str]) -> None:
    workbook = load_workbook(output, data_only=False)
    require(workbook.sheetnames == ["Municipal Summary", "Definitions"], "Unexpected worksheets")
    sheet = workbook["Municipal Summary"]
    require((sheet.max_row, sheet.max_column) == (46, 6), "Workbook dimensions changed")
    require([sheet.cell(1, column).value for column in range(1, 7)] == headers, "Headers changed")
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                require(not (isinstance(cell.value, str) and cell.value.startswith("#")), f"Cell error: {worksheet.title}!{cell.coordinate}")
    round_trip = pd.read_excel(output, sheet_name="Municipal Summary")
    require(round_trip.shape == (45, 6), "Round-trip dimensions changed")
